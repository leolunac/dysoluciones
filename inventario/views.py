from decimal import Decimal

from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.core.exceptions import ValidationError
from django.db import transaction
from django.db.models import DecimalField, ExpressionWrapper, F, Q, Sum
from django.http import HttpResponse
from django.shortcuts import redirect, render
from django.utils import timezone
from django.views.decorators.http import require_POST

from .forms import (
    ConsumoInventarioForm,
    DetalleConsumoFormSet,
    DetalleEntradaFormSet,
    EntradaInventarioForm,
    ProveedorForm,
)
from .models import (
    ConsumoInventario,
    EntradaInventario,
    ExistenciaInventario,
    MovimientoInventario,
    ProductoInventario,
)
from .permisos import exigir_consulta, exigir_gestion, puede_gestionar_inventario
from .servicios import confirmar_consumo, confirmar_entrada


@login_required
def tablero(request):
    exigir_consulta(request.user)
    buscar = request.GET.get("buscar", "").strip()
    existencias = (
        ExistenciaInventario.objects.select_related(
            "producto__accesorio",
            "ubicacion",
        )
        .filter(ubicacion__tipo="BODEGA")
        .annotate(
            valor_total=ExpressionWrapper(
                F("cantidad") * F("producto__precio_referencia"),
                output_field=DecimalField(max_digits=20, decimal_places=2),
            )
        )
        .order_by("producto__accesorio__descripcion")
    )
    if buscar:
        existencias = existencias.filter(
            Q(producto__accesorio__codigo__icontains=buscar)
            | Q(producto__accesorio__descripcion__icontains=buscar)
        )

    valor = ExpressionWrapper(
        F("cantidad") * F("producto__precio_referencia"),
        output_field=DecimalField(max_digits=20, decimal_places=2),
    )
    resumen = ExistenciaInventario.objects.filter(
        ubicacion__tipo="BODEGA",
    ).aggregate(
        unidades=Sum("cantidad"),
        valorizado=Sum(valor),
    )
    bajos = ExistenciaInventario.objects.filter(
        ubicacion__tipo="BODEGA",
        cantidad__lte=F("producto__stock_minimo"),
        producto__stock_minimo__gt=0,
    ).count()

    return render(
        request,
        "inventario/tablero.html",
        {
            "existencias": existencias[:300],
            "buscar": buscar,
            "total_unidades": resumen["unidades"] or Decimal("0.00"),
            "valor_inventario": resumen["valorizado"] or Decimal("0.00"),
            "productos_bajo_minimo": bajos,
            "puede_gestionar": puede_gestionar_inventario(request.user),
            "entradas_borrador": EntradaInventario.objects.filter(
                estado="BORRADOR"
            ).select_related("proveedor")[:10],
            "consumos_borrador": ConsumoInventario.objects.filter(
                estado="BORRADOR"
            ).select_related("cliente", "ubicacion_origen")[:10],
        },
    )


@login_required
def nueva_entrada(request):
    exigir_gestion(request.user)
    if request.method == "POST":
        form = EntradaInventarioForm(request.POST)
        formset = DetalleEntradaFormSet(request.POST)
        if form.is_valid() and formset.is_valid():
            with transaction.atomic():
                entrada = form.save(commit=False)
                entrada.creado_por = request.user
                entrada.save()
                formset.instance = entrada
                formset.save()
            messages.success(request, "Entrada guardada en borrador. Confírmela para aumentar stock.")
            return redirect("inventario:tablero")
    else:
        form = EntradaInventarioForm(initial={"fecha_factura": timezone.localdate()})
        formset = DetalleEntradaFormSet()
    return render(
        request,
        "inventario/formulario.html",
        {
            "titulo": "Nueva entrada por compra",
            "form": form,
            "formset": formset,
            "tipo_formulario": "entrada",
            "productos_disponibles": ProductoInventario.objects.filter(
                activo=True,
            ).select_related("accesorio").order_by("accesorio__descripcion"),
        },
    )


@login_required
def nuevo_proveedor(request):
    exigir_gestion(request.user)
    if request.method == "POST":
        form = ProveedorForm(request.POST)
        if form.is_valid():
            proveedor = form.save()
            messages.success(request, f"Proveedor {proveedor.nombre} creado correctamente.")
            return redirect("inventario:nueva_entrada")
    else:
        form = ProveedorForm(initial={"activo": True})
    return render(request, "inventario/proveedor_formulario.html", {"form": form})


@login_required
def nuevo_consumo(request):
    exigir_gestion(request.user)
    if request.method == "POST":
        form = ConsumoInventarioForm(request.POST)
        formset = DetalleConsumoFormSet(request.POST)
        if form.is_valid() and formset.is_valid():
            with transaction.atomic():
                consumo = form.save(commit=False)
                consumo.creado_por = request.user
                consumo.save()
                formset.instance = consumo
                formset.save()
            messages.success(request, "Consumo guardado en borrador. Confírmelo para descontar stock.")
            return redirect("inventario:tablero")
    else:
        form = ConsumoInventarioForm(initial={"fecha": timezone.localtime()})
        formset = DetalleConsumoFormSet()
    return render(
        request,
        "inventario/formulario.html",
        {
            "titulo": "Nuevo consumo por unidad",
            "form": form,
            "formset": formset,
            "tipo_formulario": "consumo",
            "productos_disponibles": ProductoInventario.objects.filter(
                activo=True,
            ).select_related("accesorio").order_by("accesorio__descripcion"),
        },
    )


@login_required
@require_POST
def confirmar_entrada_vista(request, entrada_id):
    exigir_gestion(request.user)
    try:
        confirmar_entrada(entrada_id, request.user)
        messages.success(request, "La entrada fue confirmada y el stock aumentó.")
    except (ValidationError, EntradaInventario.DoesNotExist) as exc:
        messages.error(request, "; ".join(exc.messages) if hasattr(exc, "messages") else str(exc))
    return redirect("inventario:tablero")


@login_required
@require_POST
def confirmar_consumo_vista(request, consumo_id):
    exigir_gestion(request.user)
    try:
        confirmar_consumo(consumo_id, request.user)
        messages.success(request, "El consumo fue confirmado y descontado del inventario.")
    except (ValidationError, ConsumoInventario.DoesNotExist) as exc:
        messages.error(request, "; ".join(exc.messages) if hasattr(exc, "messages") else str(exc))
    return redirect("inventario:tablero")


@login_required
def exportar_consumos_excel(request):
    exigir_consulta(request.user)
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    fecha_desde = request.GET.get("fecha_desde", "").strip()
    fecha_hasta = request.GET.get("fecha_hasta", "").strip()
    movimientos = (
        MovimientoInventario.objects.filter(tipo="CONSUMO")
        .select_related("cliente", "producto__accesorio")
        .order_by("creado", "id")
    )
    if fecha_desde:
        movimientos = movimientos.filter(creado__date__gte=fecha_desde)
    if fecha_hasta:
        movimientos = movimientos.filter(creado__date__lte=fecha_hasta)

    libro = Workbook()
    hoja = libro.active
    hoja.title = "Consumos"
    encabezados = [
        "Fecha",
        "Unidad / cliente",
        "Código",
        "Descripción",
        "Cantidad",
        "Valor unitario",
        "Valor total",
    ]
    hoja.append(encabezados)
    for celda in hoja[1]:
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill("solid", fgColor="0B5FA5")

    total_general = Decimal("0.00")
    for movimiento in movimientos.iterator():
        total = movimiento.cantidad * movimiento.precio_referencia_snapshot
        total_general += total
        hoja.append(
            [
                timezone.localtime(movimiento.creado).strftime("%Y-%m-%d %H:%M"),
                movimiento.cliente.nombre if movimiento.cliente_id else "",
                movimiento.producto.accesorio.codigo,
                movimiento.producto.accesorio.descripcion,
                float(movimiento.cantidad),
                float(movimiento.precio_referencia_snapshot),
                float(total),
            ]
        )

    hoja.append(["", "", "", "", "", "TOTAL", float(total_general)])
    hoja.cell(hoja.max_row, 6).font = Font(bold=True)
    hoja.cell(hoja.max_row, 7).font = Font(bold=True)
    for columna in (6, 7):
        for fila in range(2, hoja.max_row + 1):
            hoja.cell(fila, columna).number_format = '$#,##0.00'
    anchos = [18, 42, 16, 48, 12, 18, 18]
    for indice, ancho in enumerate(anchos, start=1):
        hoja.column_dimensions[chr(64 + indice)].width = ancho

    respuesta = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    respuesta["Content-Disposition"] = (
        f'attachment; filename="consumos_{fecha_desde or "inicio"}_{fecha_hasta or "hoy"}.xlsx"'
    )
    libro.save(respuesta)
    return respuesta
